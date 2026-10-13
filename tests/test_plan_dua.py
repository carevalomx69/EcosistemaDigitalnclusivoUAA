"""Ruta docente, paso 2: planeación con DUA y plan descargable."""
import io
import json

from docx import Document

from app import db
from app.diagnostico import IDS_ITEMS
from app.models import CursoDocente, UnidadPlan, Usuario
from app.planeacion import (DECISIONES, OPCIONES_VALIDAS, TOTAL_DECISIONES, decisiones_tomadas,
                            etiquetas_elegidas, limpiar, sugeridas)
from conftest import token_csrf


def entrar(client, correo='docente@demo.example', clave='clave-de-prueba-1'):
    client.post('/login', data={'correo': correo, 'contrasena': clave,
                                'csrf_token': token_csrf(client)})


def curso_con_diagnostico(app, correo='docente@demo.example', nombre='Programación estructurada',
                          respuestas=None):
    with app.app_context():
        u = Usuario.query.filter_by(correo=correo).one()
        r = respuestas or {i: 'si' for i in IDS_ITEMS}
        c = CursoDocente(docente_id=u.id, nombre=nombre, modalidad='hibrida', programa='ISC',
                         diagnostico_json=json.dumps(r))
        db.session.add(c)
        db.session.commit()
        return c.id


def token(client):
    return token_csrf(client, '/ruta-docente/curso/nuevo')


# ── Catálogo (sin Flask) ─────────────────────────────────────────────────
def test_catalogo_tiene_ids_unicos_y_enlaza_a_items_reales():
    ids = [d['id'] for d in DECISIONES]
    assert len(ids) == len(set(ids)) == TOTAL_DECISIONES == 8
    for d in DECISIONES:
        for _valor, _etiqueta, diag in d['opciones']:
            assert diag is None or diag in IDS_ITEMS


def test_limpiar_descarta_ids_y_valores_desconocidos():
    sucio = {'rep_formatos': ['texto', 'texto', '<x>'], 'no_existe': ['a'], 'acc_claridad': []}
    assert limpiar(sucio) == {'rep_formatos': ['texto']}


def test_decisiones_tomadas_y_sugeridas():
    assert decisiones_tomadas({'rep_formatos': ['texto'], 'acc_claridad': ['pasos']}) == 2
    assert sugeridas({'r1': 'si', 'r2': 'no', 'r3': 'no_se', 'r4': 'parcial'}) == {'r2', 'r3', 'r4'}


def test_etiquetas_elegidas_cubre_todas_las_decisiones():
    out = etiquetas_elegidas({'rep_formatos': ['video']})
    assert sum(len(filas) for _p, filas in out) == TOTAL_DECISIONES
    assert out[0][1][0][1] == ['Video con subtítulos']


# ── Rutas ────────────────────────────────────────────────────────────────
def test_plan_exige_diagnostico(client, app):
    entrar(client)
    with app.app_context():
        u = Usuario.query.one()
        c = CursoDocente(docente_id=u.id, nombre='Sin diagnóstico')
        db.session.add(c)
        db.session.commit()
        cid = c.id
    for ruta in (f'/ruta-docente/curso/{cid}/plan', f'/ruta-docente/curso/{cid}/plan/descargar'):
        r = client.get(ruta)
        assert r.status_code == 302 and r.headers['Location'].endswith('/diagnostico')


def test_flujo_agregar_editar_y_guardar_unidad(client, app):
    entrar(client)
    cid = curso_con_diagnostico(app, respuestas={**{i: 'si' for i in IDS_ITEMS}, 'r3': 'no'})
    r = client.post(f'/ruta-docente/curso/{cid}/plan/unidad',
                    data={'nombre': 'Unidad 1: Variables', 'aprendizaje': 'Declarar variables',
                          'csrf_token': token(client)})
    assert r.status_code == 302 and '/plan/unidad/' in r.headers['Location']
    ruta = r.headers['Location']
    html = client.get(ruta).get_data(as_text=True)
    assert html.count('<fieldset') == TOTAL_DECISIONES
    assert html.count('class="etiqueta-sugerida"') == 2  # video y audio, ligados a r3

    r = client.post(ruta, data={'nombre': 'Unidad 1: Variables', 'aprendizaje': '', 'notas': 'Grabar video',
                                'rep_formatos': ['texto', 'video'], 'acc_claridad': 'rubrica',
                                'csrf_token': token(client)})
    assert r.status_code == 302 and r.headers['Location'].endswith('/plan')
    with app.app_context():
        u = UnidadPlan.query.one()
        assert u.elegidas == {'rep_formatos': ['texto', 'video'], 'acc_claridad': ['rubrica']}
        assert u.notas == 'Grabar video' and u.aprendizaje is None
    html = client.get(f'/ruta-docente/curso/{cid}/plan').get_data(as_text=True)
    assert '2 de 8 decisiones tomadas' in html


def test_valores_invalidos_no_se_guardan(client, app):
    entrar(client)
    cid = curso_con_diagnostico(app)
    ruta = client.post(f'/ruta-docente/curso/{cid}/plan/unidad',
                       data={'nombre': 'Unidad X', 'csrf_token': token(client)}).headers['Location']
    client.post(ruta, data={'nombre': 'Unidad X', 'rep_formatos': ['<script>', 'texto'],
                            'inventada': 'x', 'csrf_token': token(client)})
    with app.app_context():
        assert UnidadPlan.query.one().elegidas == {'rep_formatos': ['texto']}


def test_nombre_de_unidad_obligatorio_y_sin_token_no_crea(client, app):
    entrar(client)
    cid = curso_con_diagnostico(app)
    client.post(f'/ruta-docente/curso/{cid}/plan/unidad', data={'nombre': '', 'csrf_token': token(client)})
    client.post(f'/ruta-docente/curso/{cid}/plan/unidad', data={'nombre': 'Unidad sin token'})
    with app.app_context():
        assert UnidadPlan.query.count() == 0


def test_limite_de_unidades(client, app):
    entrar(client)
    cid = curso_con_diagnostico(app)
    with app.app_context():
        db.session.add_all(UnidadPlan(curso_id=cid, orden=n, nombre=f'U{n}') for n in range(1, 21))
        db.session.commit()
    client.post(f'/ruta-docente/curso/{cid}/plan/unidad', data={'nombre': 'Una más', 'csrf_token': token(client)})
    with app.app_context():
        assert UnidadPlan.query.count() == 20


def test_unidad_y_curso_ajenos_responden_404(client, app):
    with app.app_context():
        otro = Usuario(nombre='Otro', apellidos='Docente', correo='otro@demo.example', rol='docente')
        otro.establecer_contrasena('clave-de-prueba-2')
        db.session.add(otro)
        db.session.commit()
    cid = curso_con_diagnostico(app, 'otro@demo.example')
    with app.app_context():
        u = UnidadPlan(curso_id=cid, orden=1, nombre='Ajena')
        db.session.add(u)
        db.session.commit()
        uid = u.id
    entrar(client)
    tok = token(client)
    assert client.get(f'/ruta-docente/curso/{cid}/plan').status_code == 404
    assert client.get(f'/ruta-docente/curso/{cid}/plan/unidad/{uid}').status_code == 404
    assert client.get(f'/ruta-docente/curso/{cid}/plan/descargar').status_code == 404
    assert client.post(f'/ruta-docente/curso/{cid}/plan/unidad/{uid}/eliminar',
                       data={'csrf_token': tok}).status_code == 404
    # una unidad de otro curso propio tampoco se alcanza por un id de curso distinto
    propio = curso_con_diagnostico(app)
    assert client.get(f'/ruta-docente/curso/{propio}/plan/unidad/{uid}').status_code == 404


def test_eliminar_unidad_reordena(client, app):
    entrar(client)
    cid = curso_con_diagnostico(app)
    with app.app_context():
        db.session.add_all(UnidadPlan(curso_id=cid, orden=n, nombre=f'U{n}') for n in (1, 2, 3))
        db.session.commit()
        uid = UnidadPlan.query.filter_by(nombre='U2').one().id
    r = client.post(f'/ruta-docente/curso/{cid}/plan/unidad/{uid}/eliminar', data={'csrf_token': token(client)})
    assert r.status_code == 302
    with app.app_context():
        assert [(u.nombre, u.orden) for u in UnidadPlan.query.order_by(UnidadPlan.orden)] == [('U1', 1), ('U3', 2)]


def test_eliminar_curso_borra_sus_unidades(client, app):
    entrar(client)
    cid = curso_con_diagnostico(app)
    with app.app_context():
        db.session.add(UnidadPlan(curso_id=cid, orden=1, nombre='U1'))
        db.session.commit()
    client.post(f'/ruta-docente/curso/{cid}/eliminar', data={'csrf_token': token(client)})
    with app.app_context():
        assert UnidadPlan.query.count() == 0 and CursoDocente.query.count() == 0


# ── Descarga ─────────────────────────────────────────────────────────────
def preparar_plan(client, app):
    entrar(client)
    cid = curso_con_diagnostico(app, nombre='Programación <b>estructurada</b> & más',
                                respuestas={**{i: 'si' for i in IDS_ITEMS}, 'r2': 'no'})
    with app.app_context():
        db.session.add(UnidadPlan(curso_id=cid, orden=1, nombre='Unidad 1', aprendizaje='Variables',
                                  decisiones_json=json.dumps({'rep_formatos': ['texto', 'esquema']}),
                                  notas='Revisar descripciones'))
        db.session.commit()
    return cid


def test_descarga_docx_es_un_word_valido_con_titulos_reales(client, app):
    cid = preparar_plan(client, app)
    r = client.get(f'/ruta-docente/curso/{cid}/plan/descargar?formato=docx')
    assert r.status_code == 200
    assert r.mimetype == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert 'attachment' in r.headers['Content-Disposition'] and '.docx' in r.headers['Content-Disposition']
    assert r.headers['Content-Disposition'].isascii()
    doc = Document(io.BytesIO(r.data))
    estilos = [p.style.name for p in doc.paragraphs]
    assert 'Title' in estilos and 'Heading 1' in estilos and 'Heading 2' in estilos and 'Heading 3' in estilos
    texto = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Unidad 1: Unidad 1' in texto and 'Texto con títulos y listas reales; Esquema o infografía' in texto
    assert 'Revisar descripciones' in texto and 'Sin decidir todavía.' in texto
    assert doc.core_properties.title.startswith('Plan de inclusión con DUA')
    assert doc.core_properties.language == 'es-MX'


def test_descarga_markdown(client, app):
    cid = preparar_plan(client, app)
    r = client.get(f'/ruta-docente/curso/{cid}/plan/descargar?formato=md')
    assert r.status_code == 200 and r.mimetype == 'text/markdown'
    md = r.get_data(as_text=True)
    assert md.startswith('# Plan de inclusión con DUA') and '### Unidad 1: Unidad 1' in md
    assert 'Resultado del diagnóstico' in md and 'descripción en texto' in md


def test_formato_desconocido_responde_400(client, app):
    cid = preparar_plan(client, app)
    assert client.get(f'/ruta-docente/curso/{cid}/plan/descargar?formato=exe').status_code == 400


def test_descarga_sin_unidades_no_falla(client, app):
    entrar(client)
    cid = curso_con_diagnostico(app)
    for f in ('docx', 'md'):
        assert client.get(f'/ruta-docente/curso/{cid}/plan/descargar?formato={f}').status_code == 200
